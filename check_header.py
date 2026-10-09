"""
check_header.py - cek header kolom saja (cepat), override dimension bug
"""
import sys
from openpyxl import load_workbook

filepath = sys.argv[1]
MAX_COL = 60

wb = load_workbook(filepath, read_only=True, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=MAX_COL, values_only=True)
    header = next(rows_iter)
    last_non_none = max((idx for idx, v in enumerate(header) if v is not None), default=-1)
    print(f"\n{sheet_name}: {last_non_none + 1} kolom")
    for val in header[:last_non_none+1]:
        print(f"  {val}")

wb.close()
