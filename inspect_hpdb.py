"""
inspect_hpdb.py
-----------------
Cek struktur file HPDB. File ini punya bug: tag <dimension> di XML salah
(cuma "A1"), padahal data sebenarnya punya banyak kolom & baris.
Solusi: paksa iter_rows dengan max_col/max_row eksplisit, jangan andalkan
metadata dimension.

Cara pakai:
    python inspect_hpdb.py "path/ke/file.xlsx"
"""
import sys
import os
import time
from openpyxl import load_workbook

if len(sys.argv) != 2:
    print("Cara pakai: python inspect_hpdb.py \"path/ke/file.xlsx\"")
    sys.exit(1)

filepath = sys.argv[1]
MAX_COL = 60   # lebih besar dari 53 kolom yang kita tahu
SAMPLE_ROWS = 5

print(f"File: {filepath}")
print(f"Ukuran: {os.path.getsize(filepath) / (1024*1024):.1f} MB")
print("=" * 70)

wb = load_workbook(filepath, read_only=True, data_only=True)
print(f"Sheet names: {wb.sheetnames}")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} ---")

    rows_iter = ws.iter_rows(min_row=1, max_row=SAMPLE_ROWS+1, min_col=1, max_col=MAX_COL, values_only=True)

    header = None
    for i, row in enumerate(rows_iter):
        if i == 0:
            header = row
            # Cari kolom terakhir yang tidak None
            last_non_none = max((idx for idx, v in enumerate(header) if v is not None), default=-1)
            print(f"Jumlah kolom terisi (dari {MAX_COL} dicoba): {last_non_none + 1}")
            print("\nHeader:")
            for val in header[:last_non_none+1]:
                print(f"  {val}")
        else:
            print(f"\nSample row {i}:")
            for col, val in zip(header[:last_non_none+1] if header else [], row[:last_non_none+1] if header else []):
                val_str = str(val)
                if len(val_str) > 80:
                    val_str = val_str[:80] + "...(truncated)"
                print(f"  {col} = {val_str}")

print("\n" + "=" * 70)
print("Menghitung total baris per sheet (timeout 60s per sheet)...")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    count = 0
    start = time.time()
    for row in ws.iter_rows(min_row=1, min_col=1, max_col=1, values_only=True):
        count += 1
        if time.time() - start > 60:
            print(f"  {sheet_name}: >{count:,} baris (timeout 60s, masih lanjut)")
            break
    else:
        elapsed = time.time() - start
        print(f"  {sheet_name}: {count:,} baris (termasuk header) [{elapsed:.1f}s]")

wb.close()
print("\nSelesai.")
